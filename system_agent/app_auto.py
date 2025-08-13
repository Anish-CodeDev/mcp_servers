import pyautogui
from gemini import get_rgb,extract_content,get_coordinates,get_coordinates_of_textbox,select_best_placeholder,gen_pandas_df,gen_presentation,extract_color_params
import re
import os
from reportlab.pdfgen import canvas
from docx import Document
from pptx import Presentation
from presentation import create_slide
from gemini import extract_color_params

def type_and_save(info):
    pyautogui.write(info,interval=0.1)
    pyautogui.hotkey('ctrl','s')

def select_button(name:str):
    pyautogui.screenshot('screenshots/screenshot.jpg')
    x,y = get_coordinates('screenshots/screenshot.jpg',name)
    pyautogui.click(x,y)

def execute_shortcut(shortcut:str):
    shortcut = re.sub(' ','',shortcut)
    keys = shortcut.split('+')
    print(keys)
    pyautogui.hotkey(keys[0].lower(),keys[1].lower())

def create_file_in_application(doc_type:str,filename,content):
    if doc_type == 'docx':
        doc = Document()

        doc.add_heading(content,level=1)
        doc.save(filename)
    

    elif doc_type == 'pdf':
        c = canvas.Canvas(filename)
        c.setFont("Helvetica-Bold", 16)
        c.drawString(100,750,content)
        c.save()

def type_in_textbox(task:str,data:str):
    pyautogui.screenshot('screenshots/screenshot.jpg')
    name = select_best_placeholder('screenshots/screenshot.jpg',task)
    res = get_coordinates_of_textbox('screenshots/screenshot.jpg',name)
    pyautogui.click(res[0],res[1])
    pyautogui.write(data,interval=0.1)
    pyautogui.press('Enter')

def write_to_excel(path,content):
    df = gen_pandas_df(content)
    df.to_excel(path,index=False)
    os.startfile('output.xlsx')

def create_pptx(topic:str,path:str):
    topic = extract_content(topic)
    data_dict = gen_presentation(topic)
    color = extract_color_params(topic)
    color = color.split(',')
    bg_color,title_color,content_color = color[0],color[1],color[2]
    bg_color,title_color,content_color = get_rgb(bg_color),get_rgb(title_color),get_rgb(content_color)
    slides = data_dict['slides']
    n_slides = len(slides)
    ppt = Presentation()

    for i in range(n_slides):
        create_slide(slides[i]['slide_title'],slides[i]['bullet_points'],ppt,title_color,bg_color,content_color)
    
    ppt.save(path)
    os.startfile(path)
#create_pptx('Generate a presentation on the topic AI Agents with background colour as black and the content of the color white')
